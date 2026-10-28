import io
from datetime import datetime
from typing import List, Dict, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExportService:
    @staticmethod
    def generate_attendance_dataframe(records: List[Dict[str, Any]]) -> pd.DataFrame:
        """Converts attendance list of dicts to cleaned pandas DataFrame."""
        if not records:
            return pd.DataFrame(columns=[
                "Roll Number", "Student Name", "Section", "Subject",
                "Date", "Time", "Status", "Confidence", "Session ID"
            ])

        rows = []
        for r in records:
            rows.append({
                "Roll Number": r.get("roll_number", ""),
                "Student Name": r.get("student_name", ""),
                "Section": r.get("section", ""),
                "Subject": r.get("subject", ""),
                "Date": str(r.get("date", "")),
                "Time": str(r.get("time", "")),
                "Status": str(r.get("status", "")).upper(),
                "Confidence": f"{r.get('confidence_score'):.2f}" if r.get("confidence_score") is not None else "N/A",
                "Session ID": r.get("session_id", "N/A")
            })
        return pd.DataFrame(rows)

    @staticmethod
    def generate_student_summary_dataframe(stats: List[Dict[str, Any]]) -> pd.DataFrame:
        """Converts student attendance percentage stats to pandas DataFrame."""
        if not stats:
            return pd.DataFrame(columns=[
                "Roll Number", "Student Name", "Section", "Email",
                "Total Lectures", "Present Count", "Attendance %"
            ])

        rows = []
        for s in stats:
            rows.append({
                "Roll Number": s.get("roll_number", ""),
                "Student Name": s.get("name", ""),
                "Section": s.get("section", ""),
                "Email": s.get("email", "") or "N/A",
                "Total Lectures": s.get("total_lectures", 0),
                "Present Count": s.get("present_count", 0),
                "Attendance %": f"{s.get('percentage', 0.0):.1f}%"
            })
        return pd.DataFrame(rows)

    @classmethod
    def export_to_csv(cls, df: pd.DataFrame) -> str:
        """Exports DataFrame to CSV string."""
        return df.to_csv(index=False)

    @classmethod
    def export_to_excel_bytes(
        cls,
        df: pd.DataFrame,
        sheet_name: str = "Attendance Report",
        title: str = "Classroom Attendance Report"
    ) -> bytes:
        """Exports DataFrame to formatted Excel (.xlsx) file in bytes using OpenPyXL styling."""
        output = io.BytesIO()
        
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=3)
            workbook = writer.book
            worksheet = writer.sheets[sheet_name]

            # Header styling
            title_font = Font(name="Calibri", size=16, bold=True, color="1F2937")
            meta_font = Font(name="Calibri", size=10, italic=True, color="6B7280")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
            thin_border = Border(
                left=Side(style="thin", color="E5E7EB"),
                right=Side(style="thin", color="E5E7EB"),
                top=Side(style="thin", color="E5E7EB"),
                bottom=Side(style="thin", color="E5E7EB")
            )

            # Insert Title Banner
            worksheet.cell(row=1, column=1, value=title).font = title_font
            worksheet.cell(
                row=2, column=1,
                value=f"Generated on {datetime.now().strftime('%d %b %Y at %I:%M %p')} | Facial Attendance System"
            ).font = meta_font

            # Style Table Headers (row 4)
            num_cols = len(df.columns)
            for col_num in range(1, num_cols + 1):
                cell = worksheet.cell(row=4, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Style Data Rows
            for row in range(5, len(df) + 5):
                for col in range(1, num_cols + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    # Zebra striping
                    if row % 2 == 0:
                        cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

            # Auto-fit columns
            for col in worksheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row > 2 and cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output.seek(0)
        return output.getvalue()

    @classmethod
    def export_comprehensive_excel_bytes(
        cls,
        comprehensive: Dict[str, Any],
        title: str = "Institutional Comprehensive Attendance Report"
    ) -> bytes:
        """Phase 7: Exports multi-sheet Excel workbook with Date, Subject, Semester, and Section reports."""
        output = io.BytesIO()

        # Build DataFrames
        overall = comprehensive.get("overall", {})
        df_overall = pd.DataFrame([{
            "Metric": "Total Enrolled Active Students", "Value": overall.get("total_students", 0)
        }, {
            "Metric": "Total Lecture Sessions", "Value": overall.get("total_sessions", 0)
        }, {
            "Metric": "Total Attendances Marked", "Value": overall.get("total_marked", 0)
        }, {
            "Metric": "Total Present Marks", "Value": overall.get("total_present", 0)
        }, {
            "Metric": "Total Absent Marks", "Value": overall.get("total_absent", 0)
        }, {
            "Metric": "Overall Attendance Percentage", "Value": f"{overall.get('attendance_rate', 0.0):.1f}%"
        }])

        date_wise = comprehensive.get("date_wise", [])
        df_date = pd.DataFrame([{
            "Date": d.get("date"),
            "Total Present": d.get("total_present"),
            "Total Sessions": d.get("total_sessions"),
            "Subjects Covered": ", ".join(d.get("subjects_covered", []))
        } for d in date_wise]) if date_wise else pd.DataFrame(columns=["Date", "Total Present", "Total Sessions", "Subjects Covered"])

        subj_wise = comprehensive.get("subject_wise", [])
        df_subject = pd.DataFrame([{
            "Subject": s.get("subject"),
            "Section": s.get("section"),
            "Total Sessions": s.get("total_sessions"),
            "Total Attendances": s.get("total_attendances_marked"),
            "Average per Session": s.get("average_attendance_per_session")
        } for s in subj_wise]) if subj_wise else pd.DataFrame(columns=["Subject", "Section", "Total Sessions", "Total Attendances", "Average per Session"])

        sem_wise = comprehensive.get("semester_wise", [])
        df_semester = pd.DataFrame([{
            "Semester": sem.get("semester_name"),
            "Academic Year": sem.get("academic_year_name") or "N/A",
            "Total Sessions": sem.get("total_sessions"),
            "Present Count": sem.get("present_count"),
            "Absent Count": sem.get("absent_count"),
            "Average Attendance Rate": f"{sem.get('average_attendance_rate', 0.0):.1f}%"
        } for sem in sem_wise]) if sem_wise else pd.DataFrame(columns=["Semester", "Academic Year", "Total Sessions", "Present Count", "Absent Count", "Average Attendance Rate"])

        sec_wise = comprehensive.get("section_wise", [])
        df_section = pd.DataFrame([{
            "Section": sec.get("section"),
            "Total Students": sec.get("total_students"),
            "Total Sessions": sec.get("total_sessions"),
            "Present Count": sec.get("present_count"),
            "Absent Count": sec.get("absent_count"),
            "Attendance %": f"{sec.get('attendance_percentage', 0.0):.1f}%"
        } for sec in sec_wise]) if sec_wise else pd.DataFrame(columns=["Section", "Total Students", "Total Sessions", "Present Count", "Absent Count", "Attendance %"])

        sheets = [
            ("Overall Summary", df_overall, "Institutional Overview"),
            ("Date-Wise Trends", df_date, "Date-Wise Attendance Trends"),
            ("Subject-Wise", df_subject, "Subject-Wise Performance"),
            ("Semester-Wise", df_semester, "Semester-Wise Attendance Breakdown"),
            ("Section-Wise", df_section, "Section-Wise Attendance Analysis")
        ]

        title_font = Font(name="Calibri", size=15, bold=True, color="1F2937")
        meta_font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB")
        )

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            for sname, df_sheet, stitle in sheets:
                df_sheet.to_excel(writer, sheet_name=sname, index=False, startrow=3)
                ws = writer.sheets[sname]

                # Title & meta
                ws.cell(row=1, column=1, value=f"{title} - {stitle}").font = title_font
                ws.cell(
                    row=2, column=1,
                    value=f"Generated on {datetime.now().strftime('%d %b %Y at %I:%M %p')} | Facial Attendance System"
                ).font = meta_font

                num_cols = len(df_sheet.columns)
                for col_num in range(1, num_cols + 1):
                    c = ws.cell(row=4, column=col_num)
                    c.font = header_font
                    c.fill = header_fill
                    c.alignment = Alignment(horizontal="center", vertical="center")

                for row in range(5, len(df_sheet) + 5):
                    for col in range(1, num_cols + 1):
                        c = ws.cell(row=row, column=col)
                        c.border = thin_border
                        c.alignment = Alignment(horizontal="left", vertical="center")
                        if row % 2 == 0:
                            c.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        if cell.row > 2 and cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output.seek(0)
        return output.getvalue()

export_service = ExportService()
